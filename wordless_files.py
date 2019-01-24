#
# Wordless: Files
#
# Copyright (C) 2018-2019 Ye Lei (叶磊) <blkserene@gmail.com>
#
# License Information: https://github.com/BLKSerene/Wordless/blob/master/LICENSE.txt
#

import collections
import copy
import csv
import os
import re

from PyQt5.QtCore import *
from PyQt5.QtGui import *
from PyQt5.QtWidgets import *

import bs4
import docx
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
import openpyxl
import xlrd

from wordless_checking import *
from wordless_text import *
from wordless_widgets import *
from wordless_utils import *

class Wordless_Files():
    def __init__(self, table):
        self.main = table.main
        self.table = table

    def _new_file(self, file_path):
        new_file = {}
        success_encoding_detection = True
        success_lang_detection = True
        tagged_pos = False
        tagged_non_pos = False

        new_file['selected'] = True
        new_file['path'] = os.path.normpath(file_path)
        new_file['name'], _ = os.path.splitext(os.path.basename(new_file['path']))
        new_file['name_old'] = new_file['name']

        # Detect encoding
        if self.main.settings_custom['files']['auto_detection_settings']['detect_encodings']:
            (new_file['encoding'],
             success_encoding_detection) = wordless_detection.detect_encoding(self.main, new_file['path'])
        else:
            new_file['encoding'] = self.main.settings_custom['auto_detection']['default_settings']['default_encoding']

        # Detect text type
        try:
            with open(new_file['path'], 'r', encoding = new_file['encoding']) as f:
                re_tags_pos = wordless_matching.get_re_tags(self.main, tags = 'pos')
                re_tags_non_pos = wordless_matching.get_re_tags(self.main, tags = 'non_pos')

                for line in f:
                    if re.search(re_tags_pos, line):
                        tagged_pos = True

                        break

                f.seek(0)

                for line in f:
                    if re.search(re_tags_non_pos, line):
                        tagged_non_pos = True

                        break

            if tagged_pos and tagged_non_pos:
                new_file['text_type'] = ('tokenized', 'tagged_both')
            elif tagged_pos:
                new_file['text_type'] = ('tokenized', 'tagged_pos')
            elif tagged_non_pos:
                new_file['text_type'] = ('untokenized', 'tagged_non_pos')
            else:
                new_file['text_type'] = ('untokenized', 'untagged')
        except:
            new_file['text_type'] = ('untokenized', 'untagged')

        # Detect language
        if self.main.settings_custom['files']['auto_detection_settings']['detect_langs']:
            (new_file['lang'],
             success_lang_detection) = wordless_detection.detect_lang(self.main, new_file)
        else:
            new_file['lang'] = self.main.settings_custom['auto_detection']['default_settings']['default_lang']

        return (new_file,
                success_encoding_detection,
                success_lang_detection)

    @ wordless_misc.log_timing
    def add_files(self, file_paths):
        # python-docx/Issue #276: https://github.com/python-openxml/python-docx/issues/276
        def iter_block_items(parent):
            """
            Yield each paragraph and table child within *parent*, in document order.
            Each returned value is an instance of either Table or Paragraph. *parent*
            would most commonly be a reference to a main Document object, but
            also works for a _Cell object, which itself can contain paragraphs and tables.
            """
            if isinstance(parent, Document):
                parent_elm = parent.element.body
            elif isinstance(parent, _Cell):
                parent_elm = parent._tc
            else:
                raise ValueError("something's not right")

            for child in parent_elm.iterchildren():
                if isinstance(child, CT_P):
                    yield Paragraph(child, parent)
                elif isinstance(child, CT_Tbl):
                    yield Table(child, parent)

        def iter_cell_items(parent):
            parent_elm = parent._tc

            for child in parent_elm.iterchildren():
                if isinstance(child, CT_P):
                    yield Paragraph(child, parent)
                elif isinstance(child, CT_Tbl):
                    table = Table(child, parent)

                    for row in table.rows:
                        for cell in row.cells:
                            yield from iter_cell_items(cell)

        # python-docx/Issue #40: https://github.com/python-openxml/python-docx/issues/40
        def iter_visual_cells(table):
            prior_tcs = []
            visual_cells = []

            for row in table.rows:
                visual_cells.append([])

                for cell in row.cells:
                    this_tc = cell._tc

                    if this_tc in prior_tcs:  # skip cells pointing to same `<w:tc>` element
                        continue
                    else:
                        prior_tcs.append(this_tc)

                        visual_cells[-1].append(cell)

            return visual_cells

        new_files = []

        files_encoding_detection_failed = []
        files_lang_detection_failed = []

        self.main.settings_custom['import']['files']['default_path'] = os.path.normpath(os.path.dirname(file_paths[0]))

        file_paths, files_missing = wordless_checking_file.check_files_missing(self.main, file_paths)
        file_paths, files_empty = wordless_checking_file.check_files_empty(self.main, file_paths)
        file_paths, files_duplicate = wordless_checking_file.check_files_duplicate(self.main, file_paths)
        file_paths, files_unsupported = wordless_checking_file.check_files_unsupported(self.main, file_paths)
        file_paths, files_encoding_error = wordless_checking_file.check_files_encoding_error(self.main, file_paths)

        wordless_message_box.wordless_message_box_error_files(self.main,
                                                              files_missing = files_missing,
                                                              files_empty = files_empty,
                                                              files_duplicate = files_duplicate,
                                                              files_unsupported = files_unsupported,
                                                              files_encoding_error = files_encoding_error)

        for file_path in file_paths:
            default_dir = wordless_checking_misc.check_dir(self.main.settings_custom['import']['temp_files']['default_path'])
            default_encoding = self.main.settings_custom['import']['temp_files']['default_encoding']

            file_name, file_ext = os.path.splitext(os.path.basename(file_path))
            file_ext = file_ext.lower()

            # Text Files
            if file_ext == '.txt':
                (new_file,
                 success_encoding_detection,
                 success_lang_detection) = self._new_file(file_path)

                new_files.append(new_file)

                if not success_encoding_detection:
                    files_encoding_detection_failed.append(new_file['path'])

                if not success_lang_detection:
                    files_lang_detection_failed.append(new_file['path'])
            else:
                if file_ext == ['.docx', '.xlsx', '.xls']:
                    new_path = wordless_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}.txt'))

                    # Word Documents
                    if file_ext == '.docx':
                        lines = []

                        with open(new_path, 'w', encoding = default_encoding) as f:
                            doc = docx.Document(file_path)

                            for block in iter_block_items(doc):
                                if type(block) == docx.text.paragraph.Paragraph:
                                    f.write(f'{block.text}\n')
                                elif type(block) == docx.table.Table:
                                    for row in iter_visual_cells(block):
                                        cells = []

                                        for cell in row:
                                            cells.append(' '.join([item.text for item in iter_cell_items(cell)]))

                                        f.write('\t'.join(cells) + '\n')

                    # Excel Workbooks
                    elif file_ext == '.xlsx':
                        with open(new_path, 'w', encoding = default_encoding) as f:
                            workbook = openpyxl.load_workbook(file_path, data_only = True)

                            for worksheet_name in workbook.sheetnames:
                                worksheet = workbook[worksheet_name]

                                for row in worksheet.rows:
                                    f.write('\t'.join([(cell.value if cell.value != None else '')
                                                       for cell in row]) + '\n')
                    elif file_ext == '.xls':
                        with open(new_path, 'w', encoding = default_encoding) as f:
                            workbook = xlrd.open_workbook(file_path)

                            for i_sheet in range(workbook.nsheets):
                                worksheet = workbook.sheet_by_index(i_sheet)

                                for row in range(worksheet.nrows):
                                    f.write('\t'.join([worksheet.cell_value(row, col) for col in range(worksheet.ncols)]) + '\n')

                    new_paths = [new_path]
                else:
                    # Detect encoding
                    if self.main.settings_custom['files']['auto_detection_settings']['detect_encodings']:
                        encoding_code, _ = wordless_detection.detect_encoding(self.main, file_path)
                    else:
                        encoding_code = self.main.settings_custom['encoding_detection']['default_settings']['default_encoding']

                    # CSV Files
                    if file_ext == '.csv':
                        new_path = wordless_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}.txt'))

                        with open(new_path, 'w', encoding = default_encoding) as f:
                            with open(file_path, 'r', newline = '', encoding = encoding_code) as f_csv:
                                csv_reader = csv.reader(f_csv)

                                for row in csv_reader:
                                    f.write('\t'.join(row) + '\n')

                        new_paths = [new_path]

                    # HTML Files
                    elif file_ext in ['.htm', '.html']:
                        with open(file_path, 'r', encoding = encoding_code) as f:
                            soup = bs4.BeautifulSoup(f.read(), 'lxml')

                        new_path = wordless_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}.txt'))

                        with open(new_path, 'w', encoding = default_encoding) as f:
                            f.write(soup.get_text())

                        new_paths = [new_path]

                    # Translation Memory Files
                    elif file_ext == '.tmx':
                        lines_src = []
                        lines_target = []

                        with open(file_path, 'r', encoding = encoding_code) as f:
                            soup = bs4.BeautifulSoup(f.read(), 'lxml-xml')

                            for tu in soup.find_all('tu'):
                                seg_src, seg_target = tu.find_all('seg')

                                lines_src.append(seg_src.get_text())
                                lines_target.append(seg_target.get_text())

                        path_src = wordless_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}_source.txt'))
                        path_target = wordless_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}_target.txt'))

                        with open(path_src, 'w', encoding = default_encoding) as f:
                            f.write('\n'.join(lines_src))
                            f.write('\n')

                        with open(path_target, 'w', encoding = default_encoding) as f:
                            f.write('\n'.join(lines_target))
                            f.write('\n')

                        new_paths = [path_src, path_target]

                    # Lyrics Files
                    elif file_ext == '.lrc':
                        lyrics = {}

                        with open(file_path, 'r', encoding = encoding_code) as f:
                            for line in f:
                                time_tags = []

                                line = line.strip()

                                # Strip time tags
                                while re.search(r'^\[[^\]]+?\]', line):
                                    time_tags.append(re.search(r'^\[[^\]]+?\]', line).group())

                                    line = line[len(time_tags[-1]):].strip()

                                # Strip word time tags
                                line = re.sub(r'<[^>]+?>', r'', line)
                                line = re.sub(r'\s{2,}', r' ', line).strip()

                                for time_tag in time_tags:
                                    if re.search(r'^\[[0-9]{2}:[0-5][0-9]\.[0-9]{2}\]$', time_tag):
                                        lyrics[time_tag] = line

                        new_path = wordless_checking_misc.check_new_path(f'{default_dir}{file_name}.txt')

                        with open(new_path, 'w', encoding = default_encoding) as f:
                            for _, lyrics in sorted(lyrics.items()):
                                f.write(f'{lyrics}\n')

                        new_paths = [new_path]

                    for new_path in new_paths:
                        (new_file,
                         success_encoding_detection,
                         success_lang_detection) = self._new_file(new_path)

                        new_files.append(new_file)

                        if not success_encoding_detection:
                            files_encoding_detection_failed.append(new_file['path'])

                        if not success_lang_detection:
                            files_lang_detection_failed.append(new_file['path'])

        len_files_old = len(self.main.settings_custom['files']['files_open'])

        for new_file in new_files:
            file_names = [file['name'] for file in self.main.settings_custom['files']['files_open']]
            new_file['name'] = new_file['name_old'] = wordless_checking_misc.check_new_name(new_file['name'], file_names)

            self.main.settings_custom['files']['files_open'].append(new_file)

        wordless_message_box.wordless_message_box_auto_detection_failed(self.main,
                                                                        files_encoding_detection_failed,
                                                                        files_lang_detection_failed)

        self.update_table()

        len_files_new = len(self.main.settings_custom['files']['files_open'])

        if len_files_new - len_files_old == 0:
            self.main.status_bar.showMessage('No files are newly opened!')
        elif len_files_new - len_files_old == 1:
            self.main.status_bar.showMessage('1 file has been successfully opened.')
        else:
            self.main.status_bar.showMessage(f'{len_files_new - len_files_old} files have been successfully opened.')

    def remove_files(self, indexes):
        self.main.settings_custom['files']['files_closed'].append([])

        for i in reversed(indexes):
            self.main.settings_custom['files']['files_closed'][-1].append(self.main.settings_custom['files']['files_open'].pop(i))

        self.update_table()

    def update_table(self):
        self.table.blockSignals(True)
        self.table.setUpdatesEnabled(False)

        files = self.main.settings_custom['files']['files_open']

        if files:
            self.table.clear_table(len(files))

            for i, file in enumerate(files):
                checkbox_name = QTableWidgetItem(file['name'])
                combo_box_lang = wordless_box.Wordless_Combo_Box_Lang(self.main)
                combo_box_text_type = wordless_box.Wordless_Combo_Box_Text_Type(self.main)
                combo_box_encoding = wordless_box.Wordless_Combo_Box_Encoding(self.main)

                if file['selected']:
                    checkbox_name.setCheckState(Qt.Checked)
                else:
                    checkbox_name.setCheckState(Qt.Unchecked)

                combo_box_lang.setCurrentText(wordless_conversion.to_lang_text(self.main, file['lang']))
                combo_box_text_type.setCurrentText(wordless_conversion.to_text_type_text(self.main, file['text_type']))
                combo_box_encoding.setCurrentText(wordless_conversion.to_encoding_text(self.main, file['encoding']))

                combo_box_lang.currentTextChanged.connect(lambda: self.table.itemChanged.emit(self.table.item(i, 1)))
                combo_box_text_type.currentTextChanged.connect(lambda: self.table.itemChanged.emit(self.table.item(i, 2)))
                combo_box_encoding.currentTextChanged.connect(lambda: self.table.itemChanged.emit(self.table.item(i, 4)))

                self.table.setItem(i, 0, checkbox_name)
                self.table.setCellWidget(i, 1, combo_box_lang)
                self.table.setCellWidget(i, 2, combo_box_text_type)
                self.table.setItem(i, 3, QTableWidgetItem(file['path']))
                self.table.setCellWidget(i, 4, combo_box_encoding)
        else:
            self.table.clear_table(1)

        self.table.blockSignals(False)
        self.table.setUpdatesEnabled(True)

        self.table.itemChanged.emit(self.table.item(0, 0))

    def get_selected_files(self):
        files_selected = [file for file in self.main.settings_custom['files']['files_open'] if file['selected']]

        return files_selected

    def find_file_by_name(self, file_name, selected_only = False):
        if selected_only:
            files = self.get_selected_files()
        else:
            files = self.main.settings_custom['files']['files_open']

        for file in files:
            if file['name'] == file_name:
                return file

        return None

    def find_file_by_path(self, file_path, selected_only = False):
        if selected_only:
            files = self.get_selected_files()
        else:
            files = self.main.settings_custom['files']['files_open']
            
        for file in files:
            if os.path.normcase(file['path']) == os.path.normcase(file_path):
                return file

        return None

class Wordless_Table_Files(wordless_table.Wordless_Table_Data):
    def __init__(self, main):
        super().__init__(main,
                         headers = [
                             main.tr('File Name'),
                             main.tr('Language'),
                             main.tr('Text Type'),
                             main.tr('Path'),
                             main.tr('Encoding')
                         ],
                         drag_drop_enabled = True)

        self.itemChanged.connect(self.file_item_changed)
        self.itemClicked.connect(self.file_item_changed)
        self.itemSelectionChanged.connect(self.file_selection_changed)
        self.cellDoubleClicked.connect(self.cell_double_clicked)

        self.button_open_files = QPushButton(self.tr('Add File(s)...'))
        self.button_open_dir = QPushButton(self.tr('Add Folder...'))

        self.button_close_selected = QPushButton(self.tr('Close Selected File(s)'))
        self.button_close_all = QPushButton(self.tr('Close All Files'))
        self.button_reopen = QPushButton(self.tr('Reopen Closed File(s)'))

        self.button_select_all = QPushButton(self.tr('Select All'))
        self.button_inverse = QPushButton(self.tr('Inverse'))
        self.button_deselect_all = QPushButton(self.tr('Deselect All'))

        self.button_clear.hide()
        self.button_export_selected.hide()
        self.button_export_all.hide()

        self.button_open_files.clicked.connect(self.open_files)
        self.button_open_dir.clicked.connect(lambda: self.open_dir(self.main.settings_custom['files']['subfolders']))
        self.button_close_selected.clicked.connect(self.close_selected)
        self.button_close_all.clicked.connect(self.close_all)
        self.button_reopen.clicked.connect(self.reopen)

        self.button_select_all.clicked.connect(self.select_all)
        self.button_inverse.clicked.connect(self.inverse)
        self.button_deselect_all.clicked.connect(self.deselect_all)

        self.main.__class__.open_files = self.open_files
        self.main.__class__.open_dir = self.open_dir

        self.main.__class__.close_selected = self.close_selected
        self.main.__class__.close_all = self.close_all
        self.main.__class__.reopen = self.reopen

        self.file_item_changed()

    def file_item_changed(self):
        if any([self.item(0, i) for i in range(self.columnCount())]):
            # Check duplicate file name
            for row in range(self.rowCount()):
                file_name = self.item(row, 0).text()
                file_path = self.item(row, 3).text()

                file = self.main.wordless_files.find_file_by_path(file_path)

                if file_name != file['name_old']:
                    if self.main.wordless_files.find_file_by_name(file_name):
                        self.blockSignals(True)

                        self.item(row, 0).setText(file['name_old'])
                        
                        self.blockSignals(False)

                        wordless_message_box.wordless_message_box_duplicate_file_name(self.main)

                        self.closePersistentEditor(self.item(row, 0))
                        self.editItem(self.item(row, 0))

                    break

            self.main.settings_custom['files']['files_open'].clear()

            for row in range(self.rowCount()):
                new_file = {}

                lang_text = self.cellWidget(row, 1).currentText()
                text_type_text = self.cellWidget(row, 2).currentText()
                encoding_text = self.cellWidget(row, 4).currentText()

                new_file['selected'] = True if self.item(row, 0).checkState() == Qt.Checked else False
                new_file['name'] = new_file['name_old'] = self.item(row, 0).text()
                new_file['lang'] = wordless_conversion.to_lang_code(self.main, lang_text)
                new_file['text_type'] = wordless_conversion.to_text_type_code(self.main, text_type_text)
                new_file['path'] = self.item(row, 3).text()
                new_file['encoding'] = wordless_conversion.to_encoding_code(self.main, encoding_text)

                self.main.settings_custom['files']['files_open'].append(new_file)

            self.button_select_all.setEnabled(True)
            self.button_inverse.setEnabled(True)
            self.button_deselect_all.setEnabled(True)

            self.button_close_all.setEnabled(True)
        else:
            self.button_select_all.setEnabled(False)
            self.button_inverse.setEnabled(False)
            self.button_deselect_all.setEnabled(False)

            self.button_close_all.setEnabled(False)

        if self.main.settings_custom['files']['files_closed']:
            self.button_reopen.setEnabled(True)
        else:
            self.button_reopen.setEnabled(False)

        if self.rowCount() == 0:
            self.setRowCount(1)

        self.file_selection_changed()

    def file_selection_changed(self):
        if any([self.item(0, i) for i in range(self.columnCount())]) and self.selectedIndexes():
            self.button_close_selected.setEnabled(True)
        else:
            self.button_close_selected.setEnabled(False)

    def cell_double_clicked(self, row, col):
        if col == self.find_col(self.tr('File Name')):
            self.editItem(self.item(row, col))

    def open_files(self):
        if os.path.exists(self.main.settings_custom['import']['files']['default_path']):
            default_dir = self.main.settings_custom['import']['files']['default_path']
        else:
            default_dir = self.main.settings_default['import']['files']['default_path']

        file_paths = QFileDialog.getOpenFileNames(self.main,
                                                  self.tr('Open File(s)'),
                                                  wordless_checking_misc.check_dir(default_dir),
                                                  ';;'.join(self.main.settings_global['file_types']['files']),
                                                  self.main.settings_global['file_types']['files'][-1])[0]

        if file_paths:
            self.main.wordless_files.add_files(file_paths)

    def open_dir(self, subfolders = True):
        file_paths = []

        file_dir = QFileDialog.getExistingDirectory(self.main,
                                                     self.tr('Open Folder'),
                                                     self.main.settings_custom['import']['files']['default_path'])

        if file_dir:
            if subfolders:
                for dir_path, dir_names, file_names in os.walk(file_dir):
                    for file_name in file_names:
                        file_paths.append(os.path.realpath(os.path.join(dir_path, file_name)))
            else:
                file_names = list(os.walk(file_dir))[0][2]

                for file_name in file_names:
                    file_paths.append(os.path.realpath(os.path.join(file_dir, file_name)))

            self.main.wordless_files.add_files(file_paths)

    def reopen(self):
        files = self.main.settings_custom['files']['files_closed'].pop()

        self.main.wordless_files.add_files([file['path'] for file in files])

    def select_all(self):
        if self.item(0, 0):
            for i in range(self.rowCount()):
                if self.item(i, 0).checkState() == Qt.Unchecked:
                    self.item(i, 0).setCheckState(Qt.Checked)

    def inverse(self):
        if self.item(0, 0):
            for i in range(self.rowCount()):
                if self.item(i, 0).checkState() == Qt.Checked:
                    self.item(i, 0).setCheckState(Qt.Unchecked)
                else:
                    self.item(i, 0).setCheckState(Qt.Checked)

    def deselect_all(self):
        if self.item(0, 0):
            for i in range(self.rowCount()):
                if self.item(i, 0).checkState() == Qt.Checked:
                    self.item(i, 0).setCheckState(Qt.Unchecked)

    def close_selected(self):
        self.main.wordless_files.remove_files(self.get_selected_rows())

    def close_all(self):
        self.main.wordless_files.remove_files(list(range(len(self.main.settings_custom['files']['files_open']))))

def init(main):
    def load_settings(defaults = False):
        if defaults:
            settings = copy.deepcopy(main.settings_default['files'])
        else:
            settings = copy.deepcopy(main.settings_custom['files'])

        checkbox_subfolders.setChecked(settings['folder_settings']['subfolders'])

        checkbox_detect_langs.setChecked(settings['auto_detection_settings']['detect_langs'])
        checkbox_detect_encodings.setChecked(settings['auto_detection_settings']['detect_encodings'])

        folder_settings_changed()
        auto_detection_settings_changed()

    def folder_settings_changed():
        settings = main.settings_custom['files']['folder_settings']

        settings['subfolders'] = checkbox_subfolders.isChecked()

    def auto_detection_settings_changed():
        settings = main.settings_custom['files']['auto_detection_settings']

        settings['detect_langs'] = checkbox_detect_langs.isChecked()
        settings['detect_encodings'] = checkbox_detect_encodings.isChecked()

    tab_files = wordless_layout.Wordless_Tab(main, load_settings)

    table_files = Wordless_Table_Files(main)

    tab_files.layout_table.addWidget(table_files, 0, 0, 1, 4)
    tab_files.layout_table.addWidget(table_files.button_open_files, 1, 0)
    tab_files.layout_table.addWidget(table_files.button_open_dir, 1, 1)
    tab_files.layout_table.addWidget(table_files.button_reopen, 1, 2)
    tab_files.layout_table.addWidget(table_files.button_select_all, 2, 0)
    tab_files.layout_table.addWidget(table_files.button_inverse, 2, 1)
    tab_files.layout_table.addWidget(table_files.button_deselect_all, 2, 2)
    tab_files.layout_table.addWidget(table_files.button_close_selected, 1, 3)
    tab_files.layout_table.addWidget(table_files.button_close_all, 2, 3)

    # Folder Settings
    group_box_folder_settings = QGroupBox(main.tr('Folder Settings'), main)

    checkbox_subfolders = QCheckBox(main.tr('Subfolders'), main)

    checkbox_subfolders.stateChanged.connect(folder_settings_changed)

    group_box_folder_settings.setLayout(QGridLayout())
    group_box_folder_settings.layout().addWidget(checkbox_subfolders, 0, 0)

    # Auto-detection Settings
    group_box_auto_detection_settings = QGroupBox(main.tr('Auto-detection Settings'), main)

    checkbox_detect_langs = QCheckBox(main.tr('Detect Languages'), main)
    checkbox_detect_encodings = QCheckBox(main.tr('Detect Encodings'), main)

    checkbox_detect_langs.stateChanged.connect(auto_detection_settings_changed)
    checkbox_detect_encodings.stateChanged.connect(auto_detection_settings_changed)

    group_box_auto_detection_settings.setLayout(QGridLayout())
    group_box_auto_detection_settings.layout().addWidget(checkbox_detect_langs, 0, 0)
    group_box_auto_detection_settings.layout().addWidget(checkbox_detect_encodings, 1, 0)

    tab_files.layout_settings.addWidget(group_box_folder_settings, 0, 0, Qt.AlignTop)
    tab_files.layout_settings.addWidget(group_box_auto_detection_settings, 1, 0, Qt.AlignTop)

    tab_files.layout_settings.setRowStretch(2, 1)

    load_settings()

    # Load files
    main.wordless_files = Wordless_Files(table_files)

    files = copy.deepcopy(main.settings_custom['files']['files_open'])
    file_paths = [file['path'] for file in files]

    file_paths, files_missing = wordless_checking_file.check_files_missing(main, file_paths)
    file_paths, files_empty = wordless_checking_file.check_files_empty(main, file_paths)

    wordless_message_box.wordless_message_box_error_files(main,
                                                          files_missing = files_missing,
                                                          files_empty = files_empty)

    main.settings_custom['files']['files_open'].clear()

    for file in files:
        if file['path'] in file_paths:
            main.settings_custom['files']['files_open'].append(file)

    main.wordless_files.update_table()

    return tab_files
