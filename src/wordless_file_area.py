#
# Wordless: Files
#
# Copyright (C) 2018-2019  Ye Lei (叶磊))
#
# This source file is licensed under GNU GPLv3.
# For details, see: https://github.com/BLKSerene/Wordless/blob/master/LICENSE.txt
#
# All other rights reserved.
#

import collections
import copy
import csv
import os
import re

from PyQt5.QtCore import *
from PyQt5.QtGui import *
from PyQt5.QtWidgets import *

import bs4
import docx
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
import openpyxl
import xlrd

from wordless_checking import wordless_checking_file, wordless_checking_misc
from wordless_dialogs import wordless_dialog_misc, wordless_msg_box
from wordless_utils import (wordless_conversion, wordless_detection, wordless_misc,
                            wordless_threading)
from wordless_widgets import wordless_box, wordless_layout, wordless_table

class Wordless_Worker_Add_Files(wordless_threading.Wordless_Worker):
    files_added = pyqtSignal(list, list, list, list)

    def __init__(self, main, file_paths, dialog_processed, data_received):
        super().__init__(main, dialog_processed)

        self.file_paths = file_paths

        self.files_added.connect(data_received)

    def add_files(self):
        new_files = []

        files_detection_failed_encoding = []
        files_detection_failed_text_type = []
        files_detection_failed_lang = []

        if self.file_paths:
            len_file_paths = len(self.file_paths)
            self.main.settings_custom['import']['files']['default_path'] = os.path.normpath(os.path.dirname(self.file_paths[0]))

            for i, file_path in enumerate(self.file_paths):
                self.progress_updated.emit(self.tr(f'Loading files ... ({i + 1}/{len_file_paths})'))

                default_dir = wordless_checking_misc.check_dir(self.main.settings_custom['import']['temp_files']['default_path'])
                default_encoding = self.main.settings_custom['import']['temp_files']['default_encoding']

                file_name, file_ext = os.path.splitext(os.path.basename(file_path))
                file_ext = file_ext.lower()

                # Text Files
                if file_ext == '.txt':
                    (new_file,
                     detection_success_encoding,
                     detection_success_text_type,
                     detection_success_lang) = self.main.wordless_files._new_file(file_path)

                    new_files.append(new_file)

                    if not detection_success_encoding:
                        files_detection_failed_encoding.append(new_file['path'])

                    if not detection_success_text_type:
                        files_detection_failed_text_type.append(new_file['path'])

                    if not detection_success_lang:
                        files_detection_failed_lang.append(new_file['path'])
                else:
                    if file_ext == ['.docx', '.xlsx', '.xls']:
                        new_path = wordless_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}.txt'))

                        # Word Documents
                        if file_ext == '.docx':
                            lines = []

                            with open(new_path, 'w', encoding = default_encoding) as f:
                                doc = docx.Document(file_path)

                                for block in self.iter_block_items(doc):
                                    if type(block) == docx.text.paragraph.Paragraph:
                                        f.write(f'{block.text}\n')
                                    elif type(block) == docx.table.Table:
                                        for row in self.iter_visual_cells(block):
                                            cells = []

                                            for cell in row:
                                                cells.append(' '.join([item.text for item in self.iter_cell_items(cell)]))

                                            f.write('\t'.join(cells) + '\n')

                        # Excel Workbooks
                        elif file_ext == '.xlsx':
                            with open(new_path, 'w', encoding = default_encoding) as f:
                                workbook = openpyxl.load_workbook(file_path, data_only = True)

                                for worksheet_name in workbook.sheetnames:
                                    worksheet = workbook[worksheet_name]

                                    for row in worksheet.rows:
                                        f.write('\t'.join([(cell.value if cell.value != None else '')
                                                           for cell in row]) + '\n')
                        elif file_ext == '.xls':
                            with open(new_path, 'w', encoding = default_encoding) as f:
                                workbook = xlrd.open_workbook(file_path)

                                for i_sheet in range(workbook.nsheets):
                                    worksheet = workbook.sheet_by_index(i_sheet)

                                    for row in range(worksheet.nrows):
                                        f.write('\t'.join([worksheet.cell_value(row, col) for col in range(worksheet.ncols)]) + '\n')

                        new_paths = [new_path]
                    else:
                        # Detect encoding
                        if self.main.settings_custom['files']['auto_detection_settings']['detect_encodings']:
                            encoding_code, _ = wordless_detection.detect_encoding(self.main, file_path)
                        else:
                            encoding_code = self.main.settings_custom['encoding_detection']['default_settings']['default_encoding']

                        # CSV Files
                        if file_ext == '.csv':
                            new_path = wordless_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}.txt'))

                            with open(new_path, 'w', encoding = default_encoding) as f:
                                with open(file_path, 'r', newline = '', encoding = encoding_code) as f_csv:
                                    csv_reader = csv.reader(f_csv)

                                    for row in csv_reader:
                                        f.write('\t'.join(row) + '\n')

                            new_paths = [new_path]

                        # HTML Files
                        elif file_ext in ['.htm', '.html']:
                            with open(file_path, 'r', encoding = encoding_code) as f:
                                soup = bs4.BeautifulSoup(f.read(), 'lxml')

                            new_path = wordless_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}.txt'))

                            with open(new_path, 'w', encoding = default_encoding) as f:
                                f.write(soup.get_text())

                            new_paths = [new_path]

                        # Translation Memory Files
                        elif file_ext == '.tmx':
                            lines_src = []
                            lines_target = []

                            with open(file_path, 'r', encoding = encoding_code) as f:
                                soup = bs4.BeautifulSoup(f.read(), 'lxml-xml')

                                for tu in soup.find_all('tu'):
                                    seg_src, seg_target = tu.find_all('seg')

                                    lines_src.append(seg_src.get_text())
                                    lines_target.append(seg_target.get_text())

                            path_src = wordless_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}_source.txt'))
                            path_target = wordless_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}_target.txt'))

                            with open(path_src, 'w', encoding = default_encoding) as f:
                                f.write('\n'.join(lines_src))
                                f.write('\n')

                            with open(path_target, 'w', encoding = default_encoding) as f:
                                f.write('\n'.join(lines_target))
                                f.write('\n')

                            new_paths = [path_src, path_target]

                        # Lyrics Files
                        elif file_ext == '.lrc':
                            lyrics = {}

                            with open(file_path, 'r', encoding = encoding_code) as f:
                                for line in f:
                                    time_tags = []

                                    line = line.strip()

                                    # Strip time tags
                                    while re.search(r'^\[[^\]]+?\]', line):
                                        time_tags.append(re.search(r'^\[[^\]]+?\]', line).group())

                                        line = line[len(time_tags[-1]):].strip()

                                    # Strip word time tags
                                    line = re.sub(r'<[^>]+?>', r'', line)
                                    line = re.sub(r'\s{2,}', r' ', line).strip()

                                    for time_tag in time_tags:
                                        if re.search(r'^\[[0-9]{2}:[0-5][0-9]\.[0-9]{2}\]$', time_tag):
                                            lyrics[time_tag] = line

                            new_path = wordless_checking_misc.check_new_path(f'{default_dir}{file_name}.txt')

                            with open(new_path, 'w', encoding = default_encoding) as f:
                                for _, lyrics in sorted(lyrics.items()):
                                    f.write(f'{lyrics}\n')

                            new_paths = [new_path]

                        for new_path in new_paths:
                            (new_file,
                             detection_success_encoding,
                             detection_success_text_type,
                             detection_success_lang) = self.main.wordless_files._new_file(new_path)

                            new_files.append(new_file)

                            if not detection_success_encoding:
                                files_detection_failed_encoding.append(new_file['path'])

                            if not detection_success_text_type:
                                files_detection_failed_text_type.append(new_file['path'])

                            if not detection_success_lang:
                                files_detection_failed_lang.append(new_file['path'])

        self.files_added.emit(new_files,
                              files_detection_failed_encoding,
                              files_detection_failed_text_type,
                              files_detection_failed_lang)

    # python-docx/Issue #276: https://github.com/python-openxml/python-docx/issues/276
    def iter_block_items(self, parent):
        """
        Yield each paragraph and table child within *parent*, in document order.
        Each returned value is an instance of either Table or Paragraph. *parent*
        would most commonly be a reference to a main Document object, but
        also works for a _Cell object, which itself can contain paragraphs and tables.
        """
        if isinstance(parent, Document):
            parent_elm = parent.element.body
        elif isinstance(parent, _Cell):
            parent_elm = parent._tc
        else:
            raise ValueError("something's not right")

        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    def iter_cell_items(self, parent):
        parent_elm = parent._tc

        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                table = Table(child, parent)

                for row in table.rows:
                    for cell in row.cells:
                        yield from self.iter_cell_items(cell)

    # python-docx/Issue #40: https://github.com/python-openxml/python-docx/issues/40
    def iter_visual_cells(self, table):
        prior_tcs = []
        visual_cells = []

        for row in table.rows:
            visual_cells.append([])

            for cell in row.cells:
                this_tc = cell._tc

                if this_tc in prior_tcs:  # skip cells pointing to same `<w:tc>` element
                    continue
                else:
                    prior_tcs.append(this_tc)

                    visual_cells[-1].append(cell)

        return visual_cells

class Wordless_Files():
    def __init__(self, table):
        self.main = table.main
        self.table = table

    def _new_file(self, file_path):
        new_file = {}

        detection_success_encoding = True
        detection_success_text_type = True
        detection_success_lang = True

        new_file['selected'] = True
        new_file['path'] = os.path.normpath(file_path)
        new_file['name'], _ = os.path.splitext(os.path.basename(new_file['path']))
        new_file['name_old'] = new_file['name']

        # Detect encodings
        if self.main.settings_custom['files']['auto_detection_settings']['detect_encodings']:
            (new_file['encoding'],
             detection_success_encoding) = wordless_detection.detect_encoding(self.main, new_file['path'])
        else:
            new_file['encoding'] = self.main.settings_custom['auto_detection']['default_settings']['default_encoding']

        # Detect text types
        if self.main.settings_custom['files']['auto_detection_settings']['detect_text_types']:
            (new_file['text_type'],
             detection_success_text_type) = wordless_detection.detect_text_type(self.main, new_file)
        else:
            new_file['text_type'] = self.main.settings_custom['auto_detection']['default_settings']['default_text_type']

        # Detect languages
        if self.main.settings_custom['files']['auto_detection_settings']['detect_langs']:
            (new_file['lang'],
             detection_success_lang) = wordless_detection.detect_lang(self.main, new_file)
        else:
            new_file['lang'] = self.main.settings_custom['auto_detection']['default_settings']['default_lang']

        return (new_file,
                detection_success_encoding,
                detection_success_text_type,
                detection_success_lang)

    @wordless_misc.log_timing
    def add_files(self, file_paths):
        def data_received(new_files,
                          files_detection_failed_encoding,
                          files_detection_failed_text_type,
                          files_detection_failed_lang):
            len_files_old = len(self.main.settings_custom['files']['files_open'])

            for new_file in new_files:
                file_names = [file['name'] for file in self.main.settings_custom['files']['files_open']]
                new_file['name'] = new_file['name_old'] = wordless_checking_misc.check_new_name(new_file['name'], file_names)

                self.main.settings_custom['files']['files_open'].append(new_file)

            wordless_msg_box.wordless_msg_box_detection_failed(self.main,
                                                               files_detection_failed_encoding,
                                                               files_detection_failed_text_type,
                                                               files_detection_failed_lang)

            self.update_table()

            len_files_new = len(self.main.settings_custom['files']['files_open'])

            if len_files_new - len_files_old == 0:
                self.main.statusBar().showMessage('No files are newly opened!')
            elif len_files_new - len_files_old == 1:
                self.main.statusBar().showMessage('1 file has been successfully opened.')
            else:
                self.main.statusBar().showMessage(f'{len_files_new - len_files_old} files have been successfully opened.')

            dialog_progress.accept()

        file_paths, files_empty = wordless_checking_file.check_files_empty(self.main, file_paths)
        file_paths, files_duplicate = wordless_checking_file.check_files_duplicate(self.main, file_paths)
        file_paths, files_unsupported = wordless_checking_file.check_files_unsupported(self.main, file_paths)
        file_paths, files_parsing_error = wordless_checking_file.check_files_parsing_error(self.main, file_paths)

        wordless_msg_box.wordless_msg_box_file_error_on_opening(self.main,
                                                                files_empty = files_empty,
                                                                files_duplicate = files_duplicate,
                                                                files_unsupported = files_unsupported,
                                                                files_parsing_error = files_parsing_error)

        dialog_progress = wordless_dialog_misc.Wordless_Dialog_Progress_Add_Files(self.main)

        worker_add_files = Wordless_Worker_Add_Files(self.main, file_paths, dialog_progress, data_received)
        thread_add_files = wordless_threading.Wordless_Thread_Add_Files(worker_add_files)

        thread_add_files.start()

        dialog_progress.exec_()

        thread_add_files.quit()
        thread_add_files.wait()

    def remove_files(self, indexes):
        self.main.settings_custom['files']['files_closed'].append([])

        for i in reversed(indexes):
            self.main.settings_custom['files']['files_closed'][-1].append(self.main.settings_custom['files']['files_open'].pop(i))

        self.update_table()

    def update_table(self):
        self.table.blockSignals(True)
        self.table.setUpdatesEnabled(False)

        files = self.main.settings_custom['files']['files_open']

        if files:
            self.table.clear_table(len(files))

            for i, file in enumerate(files):
                checkbox_name = QTableWidgetItem(file['name'])
                combo_box_lang = wordless_box.Wordless_Combo_Box_Lang(self.table)
                combo_box_text_type = wordless_box.Wordless_Combo_Box_Text_Type(self.table)
                combo_box_encoding = wordless_box.Wordless_Combo_Box_Encoding(self.table)

                if file['selected']:
                    checkbox_name.setCheckState(Qt.Checked)
                else:
                    checkbox_name.setCheckState(Qt.Unchecked)

                combo_box_lang.setCurrentText(wordless_conversion.to_lang_text(self.main, file['lang']))
                combo_box_text_type.setCurrentText(wordless_conversion.to_text_type_text(self.main, file['text_type']))
                combo_box_encoding.setCurrentText(wordless_conversion.to_encoding_text(self.main, file['encoding']))

                combo_box_lang.currentTextChanged.connect(lambda: self.table.itemChanged.emit(self.table.item(i, 1)))
                combo_box_text_type.currentTextChanged.connect(lambda: self.table.itemChanged.emit(self.table.item(i, 2)))
                combo_box_encoding.currentTextChanged.connect(lambda: self.table.itemChanged.emit(self.table.item(i, 4)))

                self.table.setItem(i, 0, checkbox_name)
                self.table.setCellWidget(i, 1, combo_box_lang)
                self.table.setCellWidget(i, 2, combo_box_text_type)
                self.table.setItem(i, 3, QTableWidgetItem(file['path']))
                self.table.setCellWidget(i, 4, combo_box_encoding)
        else:
            self.table.clear_table(1)

        self.table.blockSignals(False)
        self.table.setUpdatesEnabled(True)

        self.table.itemChanged.emit(self.table.item(0, 0))

    def get_selected_files(self):
        files_selected = [file for file in self.main.settings_custom['files']['files_open'] if file['selected']]

        return files_selected

    def find_file_by_name(self, file_name, selected_only = False):
        if selected_only:
            files = self.get_selected_files()
        else:
            files = self.main.settings_custom['files']['files_open']

        for file in files:
            if file['name'] == file_name:
                return file

        return None

    def find_file_by_path(self, file_path, selected_only = False):
        if selected_only:
            files = self.get_selected_files()
        else:
            files = self.main.settings_custom['files']['files_open']
            
        for file in files:
            if os.path.normcase(file['path']) == os.path.normcase(file_path):
                return file

        return None

class Wordless_Table_Files(wordless_table.Wordless_Table):
    def __init__(self, parent):
        super().__init__(parent,
                         headers = [
                             parent.tr('File Name'),
                             parent.tr('Language'),
                             parent.tr('Text Type'),
                             parent.tr('Path'),
                             parent.tr('Encoding')
                         ],
                         drag_drop_enabled = True)

        self.itemChanged.connect(self.file_item_changed)
        self.itemClicked.connect(self.file_item_changed)
        self.itemSelectionChanged.connect(self.file_selection_changed)
        self.cellDoubleClicked.connect(self.cell_double_clicked)

        self.button_open_files = QPushButton(self.tr('Add File(s)...'))
        self.button_open_dir = QPushButton(self.tr('Add Folder...'))
        self.button_reopen = QPushButton(self.tr('Reopen Closed File(s)'))

        self.button_select_all = QPushButton(self.tr('Select All'))
        self.button_invert_selection = QPushButton(self.tr('Invert Selection'))
        self.button_deselect_all = QPushButton(self.tr('Deselect All'))

        self.button_close_selected = QPushButton(self.tr('Close Selected'))
        self.button_close_all = QPushButton(self.tr('Close All'))

        self.button_open_files.clicked.connect(self.open_files)
        self.button_open_dir.clicked.connect(self.open_dir)
        self.button_reopen.clicked.connect(self.reopen)

        self.button_select_all.clicked.connect(self.select_all)
        self.button_invert_selection.clicked.connect(self.invert_selection)
        self.button_deselect_all.clicked.connect(self.deselect_all)

        self.button_close_selected.clicked.connect(self.close_selected)
        self.button_close_all.clicked.connect(self.close_all)

        # Menu
        self.main.find_menu_item(self.tr('Open File(s)...')).triggered.connect(self.open_files)
        self.main.find_menu_item(self.tr('Open Folder...')).triggered.connect(self.open_dir)

        self.main.find_menu_item(self.tr('Reopen Closed File(s)')).triggered.connect(self.reopen)

        self.main.find_menu_item(self.tr('Select All')).triggered.connect(self.select_all)
        self.main.find_menu_item(self.tr('Invert Selection')).triggered.connect(self.invert_selection)
        self.main.find_menu_item(self.tr('Deselect All')).triggered.connect(self.deselect_all)

        self.main.find_menu_item(self.tr('Close Selected')).triggered.connect(self.close_selected)
        self.main.find_menu_item(self.tr('Close All')).triggered.connect(self.close_all)

        self.file_item_changed()

    def file_item_changed(self):
        if any([self.item(0, i) for i in range(self.columnCount())]):
            # Check duplicate file name
            for row in range(self.rowCount()):
                file_name = self.item(row, 0).text()
                file_path = self.item(row, 3).text()

                file = self.main.wordless_files.find_file_by_path(file_path)

                if file_name != file['name_old']:
                    if self.main.wordless_files.find_file_by_name(file_name):
                        self.blockSignals(True)

                        self.item(row, 0).setText(file['name_old'])
                        
                        self.blockSignals(False)

                        wordless_msg_box.wordless_msg_box_duplicate_file_name(self.main)

                        self.closePersistentEditor(self.item(row, 0))
                        self.editItem(self.item(row, 0))

                    break

            self.main.settings_custom['files']['files_open'].clear()

            for row in range(self.rowCount()):
                new_file = {}

                lang_text = self.cellWidget(row, 1).currentText()
                text_type_text = self.cellWidget(row, 2).currentText()
                encoding_text = self.cellWidget(row, 4).currentText()

                new_file['selected'] = True if self.item(row, 0).checkState() == Qt.Checked else False
                new_file['name'] = new_file['name_old'] = self.item(row, 0).text()
                new_file['lang'] = wordless_conversion.to_lang_code(self.main, lang_text)
                new_file['text_type'] = wordless_conversion.to_text_type_code(self.main, text_type_text)
                new_file['path'] = self.item(row, 3).text()
                new_file['encoding'] = wordless_conversion.to_encoding_code(self.main, encoding_text)

                self.main.settings_custom['files']['files_open'].append(new_file)

            self.button_select_all.setEnabled(True)
            self.button_invert_selection.setEnabled(True)
            self.button_deselect_all.setEnabled(True)

            self.button_close_all.setEnabled(True)
        else:
            self.button_select_all.setEnabled(False)
            self.button_invert_selection.setEnabled(False)
            self.button_deselect_all.setEnabled(False)

            self.button_close_all.setEnabled(False)

        if self.main.settings_custom['files']['files_closed']:
            self.button_reopen.setEnabled(True)
        else:
            self.button_reopen.setEnabled(False)

        # Menu
        if any([self.item(0, i) for i in range(self.columnCount())]):
            self.main.find_menu_item(self.tr('Select All')).setEnabled(True)
            self.main.find_menu_item(self.tr('Invert Selection')).setEnabled(True)
            self.main.find_menu_item(self.tr('Deselect All')).setEnabled(True)

            self.main.find_menu_item(self.tr('Close All')).setEnabled(True)
        else:
            self.main.find_menu_item(self.tr('Select All')).setEnabled(False)
            self.main.find_menu_item(self.tr('Invert Selection')).setEnabled(False)
            self.main.find_menu_item(self.tr('Deselect All')).setEnabled(False)

            self.main.find_menu_item(self.tr('Close All')).setEnabled(False)

        if self.main.settings_custom['files']['files_closed']:
            self.main.find_menu_item(self.tr('Reopen Closed File(s)')).setEnabled(True)
        else:
            self.main.find_menu_item(self.tr('Reopen Closed File(s)')).setEnabled(False)

        if self.rowCount() == 0:
            self.setRowCount(1)

        self.file_selection_changed()

    def file_selection_changed(self):
        if any([self.item(0, i) for i in range(self.columnCount())]) and self.selectedIndexes():
            self.button_close_selected.setEnabled(True)
        else:
            self.button_close_selected.setEnabled(False)

        # Menu
        if any([self.item(0, i) for i in range(self.columnCount())]) and self.selectedIndexes():
            self.main.find_menu_item(self.tr('Close Selected')).setEnabled(True)
        else:
            self.main.find_menu_item(self.tr('Close Selected')).setEnabled(False)

    def cell_double_clicked(self, row, col):
        if col == self.find_col(self.tr('File Name')):
            self.editItem(self.item(row, col))

    def open_files(self):
        if os.path.exists(self.main.settings_custom['import']['files']['default_path']):
            default_dir = self.main.settings_custom['import']['files']['default_path']
        else:
            default_dir = self.main.settings_default['import']['files']['default_path']

        file_paths = QFileDialog.getOpenFileNames(self.main,
                                                  self.tr('Open File(s)'),
                                                  wordless_checking_misc.check_dir(default_dir),
                                                  ';;'.join(self.main.settings_global['file_types']['files']),
                                                  self.main.settings_global['file_types']['files'][-1])[0]

        if file_paths:
            self.main.wordless_files.add_files(file_paths)

    def open_dir(self):
        file_paths = []

        file_dir = QFileDialog.getExistingDirectory(self.main,
                                                    self.tr('Open Folder'),
                                                    self.main.settings_custom['import']['files']['default_path'])

        if file_dir:
            if self.main.settings_custom['files']['subfolders']:
                for dir_path, dir_names, file_names in os.walk(file_dir):
                    for file_name in file_names:
                        file_paths.append(os.path.realpath(os.path.join(dir_path, file_name)))
            else:
                file_names = list(os.walk(file_dir))[0][2]

                for file_name in file_names:
                    file_paths.append(os.path.realpath(os.path.join(file_dir, file_name)))

            self.main.wordless_files.add_files(file_paths)

    def reopen(self):
        files = self.main.settings_custom['files']['files_closed'].pop()

        self.main.wordless_files.add_files([file['path'] for file in files])

    def select_all(self):
        if self.item(0, 0):
            for i in range(self.rowCount()):
                if self.item(i, 0).checkState() == Qt.Unchecked:
                    self.item(i, 0).setCheckState(Qt.Checked)

    def invert_selection(self):
        if self.item(0, 0):
            for i in range(self.rowCount()):
                if self.item(i, 0).checkState() == Qt.Checked:
                    self.item(i, 0).setCheckState(Qt.Unchecked)
                else:
                    self.item(i, 0).setCheckState(Qt.Checked)

    def deselect_all(self):
        if self.item(0, 0):
            for i in range(self.rowCount()):
                if self.item(i, 0).checkState() == Qt.Checked:
                    self.item(i, 0).setCheckState(Qt.Unchecked)

    def close_selected(self):
        self.main.wordless_files.remove_files(self.get_selected_rows())

    def close_all(self):
        self.main.wordless_files.remove_files(list(range(len(self.main.settings_custom['files']['files_open']))))

class Wrapper_File_Area(wordless_layout.Wordless_Wrapper_File_Area):
    def __init__(self, main):
        super().__init__(main)

        # Table
        self.table_files = Wordless_Table_Files(self)

        self.wrapper_table.layout().addWidget(self.table_files, 0, 0, 1, 4)
        self.wrapper_table.layout().addWidget(self.table_files.button_open_files, 1, 0)
        self.wrapper_table.layout().addWidget(self.table_files.button_open_dir, 1, 1)
        self.wrapper_table.layout().addWidget(self.table_files.button_reopen, 1, 2)
        self.wrapper_table.layout().addWidget(self.table_files.button_select_all, 2, 0)
        self.wrapper_table.layout().addWidget(self.table_files.button_invert_selection, 2, 1)
        self.wrapper_table.layout().addWidget(self.table_files.button_deselect_all, 2, 2)
        self.wrapper_table.layout().addWidget(self.table_files.button_close_selected, 1, 3)
        self.wrapper_table.layout().addWidget(self.table_files.button_close_all, 2, 3)

        # Folder Settings
        self.group_box_folder_settings = QGroupBox(self.tr('Folder Settings'), self)

        self.checkbox_subfolders = QCheckBox(self.tr('Subfolders'), self)

        self.checkbox_subfolders.stateChanged.connect(self.folder_settings_changed)

        self.group_box_folder_settings.setLayout(wordless_layout.Wordless_Layout())
        self.group_box_folder_settings.layout().addWidget(self.checkbox_subfolders, 0, 0)

        # Auto-detection Settings
        self.group_box_auto_detection_settings = QGroupBox(self.tr('Auto-detection Settings'), self)

        self.checkbox_detect_langs = QCheckBox(self.tr('Detect Languages'), self)
        self.checkbox_detect_text_types = QCheckBox(self.tr('Detect Text Types'), self)
        self.checkbox_detect_encodings = QCheckBox(self.tr('Detect Encodings'), self)

        self.checkbox_detect_langs.stateChanged.connect(self.auto_detection_settings_changed)
        self.checkbox_detect_text_types.stateChanged.connect(self.auto_detection_settings_changed)
        self.checkbox_detect_encodings.stateChanged.connect(self.auto_detection_settings_changed)

        self.group_box_auto_detection_settings.setLayout(wordless_layout.Wordless_Layout())
        self.group_box_auto_detection_settings.layout().addWidget(self.checkbox_detect_langs, 0, 0)
        self.group_box_auto_detection_settings.layout().addWidget(self.checkbox_detect_text_types, 0, 1)
        self.group_box_auto_detection_settings.layout().addWidget(self.checkbox_detect_encodings, 1, 0)

        self.wrapper_settings.layout().addWidget(self.group_box_folder_settings, 0, 0)
        self.wrapper_settings.layout().addWidget(self.group_box_auto_detection_settings, 1, 0)

        self.wrapper_settings.layout().setRowStretch(2, 1)

        self.load_settings()

        # Load files
        self.main.wordless_files = Wordless_Files(self.table_files)

        files = copy.deepcopy(self.main.settings_custom['files']['files_open'])
        file_paths = [file['path'] for file in files]

        file_paths, files_missing = wordless_checking_file.check_files_missing(self.main, file_paths)
        file_paths, files_empty = wordless_checking_file.check_files_empty(self.main, file_paths)

        wordless_msg_box.wordless_msg_box_file_error_on_startup(self.main,
                                                                files_missing = files_missing,
                                                                files_empty = files_empty)

        self.main.settings_custom['files']['files_open'].clear()

        for file in files:
            if file['path'] in file_paths:
                self.main.settings_custom['files']['files_open'].append(file)

        self.main.wordless_files.update_table()

    def load_settings(self, defaults = False):
        if defaults:
            settings = copy.deepcopy(self.main.settings_default['files'])
        else:
            settings = copy.deepcopy(self.main.settings_custom['files'])

        self.checkbox_subfolders.setChecked(settings['folder_settings']['subfolders'])

        self.checkbox_detect_langs.setChecked(settings['auto_detection_settings']['detect_langs'])
        self.checkbox_detect_text_types.setChecked(settings['auto_detection_settings']['detect_text_types'])
        self.checkbox_detect_encodings.setChecked(settings['auto_detection_settings']['detect_encodings'])

        self.folder_settings_changed()
        self.auto_detection_settings_changed()

    def folder_settings_changed(self):
        settings = self.main.settings_custom['files']['folder_settings']

        settings['subfolders'] = self.checkbox_subfolders.isChecked()

    def auto_detection_settings_changed(self):
        settings = self.main.settings_custom['files']['auto_detection_settings']

        settings['detect_langs'] = self.checkbox_detect_langs.isChecked()
        settings['detect_text_types'] = self.checkbox_detect_text_types.isChecked()
        settings['detect_encodings'] = self.checkbox_detect_encodings.isChecked()
