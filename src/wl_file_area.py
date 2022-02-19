# ----------------------------------------------------------------------
# Wordless: Files
# Copyright (C) 2018-2022  Ye Lei (叶磊)
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <http://www.gnu.org/licenses/>.
# ----------------------------------------------------------------------

import copy
import csv
import os
import re
import traceback

import bs4
import docx
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
import openpyxl
from PyQt5.QtCore import *
from PyQt5.QtGui import *
from PyQt5.QtWidgets import *

from wl_checking import wl_checking_files, wl_checking_misc
from wl_dialogs import wl_dialogs, wl_dialogs_errs, wl_dialogs_misc, wl_msg_boxes
from wl_nlp import wl_matching, wl_texts
from wl_utils import wl_conversion, wl_detection, wl_misc, wl_msgs, wl_threading
from wl_widgets import wl_boxes, wl_buttons, wl_item_delegates, wl_layouts, wl_tables

class Wl_Worker_Add_Files(wl_threading.Wl_Worker):
    worker_done = pyqtSignal(str, list)

    def run(self):
        err_msg = ''
        new_files = []

        try:
            len_file_paths = len(self.file_paths)

            for i, file_path in enumerate(self.file_paths):
                self.progress_updated.emit(self.tr(f'Adding files... ({i + 1}/{len_file_paths})'))

                file_path = wl_misc.get_normalized_path(file_path)
                file_name, file_ext = os.path.splitext(os.path.basename(file_path))
                file_ext = file_ext.lower()

                new_file = {'selected': True, 'path_original': file_path}

                # Check for duplicate file names
                file_names = [
                    *self.main.wl_file_area.get_selected_file_names(),
                    *[file['name'] for file in self.table.files_to_open],
                    *[new_file['name'] for new_file in new_files]
                ]

                new_file['name'] = new_file['name_old'] = wl_checking_misc.check_new_name(file_name, file_names)

                # Path, Tokenized, Tagged
                default_dir = wl_checking_misc.check_dir(self.main.settings_custom['imp']['temp_files']['default_path'])

                if file_ext == '.xml':
                    new_file['path'] = os.path.join(default_dir, f'{file_name}.xml')
                    new_file['tokenized'] = 'Yes'
                    new_file['tagged'] = 'Yes'
                else:
                    new_file['path'] = os.path.join(default_dir, f'{file_name}.txt')
                    new_file['tokenized'] = self.main.settings_custom['files']['default_settings']['tokenized']
                    new_file['tagged'] = self.main.settings_custom['files']['default_settings']['tagged']

                # Check for duplicate files
                new_file['path'] = wl_checking_misc.check_new_path(new_file['path'])

                # Detect encodings
                default_encoding = self.main.settings_custom['files']['default_settings']['encoding']

                if file_ext in ['.docx', '.xlsx']:
                    new_file['encoding'] = default_encoding
                else:
                    if self.main.settings_custom['file_area']['dialog_open_files']['auto_detect_encodings']:
                        new_file['encoding'] = wl_detection.detect_encoding(self.main, file_path)
                    else:
                        new_file['encoding'] = default_encoding

                # Cleanse contents before language detection
                if file_ext != '.tmx':
                    if file_ext in ['.txt', '.xml']:
                        with open(file_path, 'r', encoding = new_file['encoding'], errors = 'replace') as f:
                            new_file['text'] = f.read()
                    # CSV files
                    elif file_ext == '.csv':
                        lines = []

                        with open(file_path, 'r', encoding = new_file['encoding'], errors = 'replace', newline = '') as f:
                            # Remove NULL bytes to avoid error
                            csv_reader = csv.reader([line.replace('\0', '') for line in f])

                            for row in csv_reader:
                                lines.append('\t'.join(row))

                        new_file['text'] = '\n'.join(lines)
                    # HTML pages
                    elif file_ext in ['.htm', '.html']:
                        with open(file_path, 'r', encoding = new_file['encoding'], errors = 'replace') as f:
                            soup = bs4.BeautifulSoup(f.read(), 'lxml')

                        new_file['text'] = soup.get_text()
                    # Microsoft Word documents
                    elif file_ext == '.docx':
                        lines = []
                        doc = docx.Document(file_path)

                        for block in self.iter_block_items(doc):
                            if type(block) == docx.text.paragraph.Paragraph:
                                lines.append(block.text)
                            elif type(block) == docx.table.Table:
                                for row in self.iter_visual_cells(block):
                                    cells = []

                                    for cell in row:
                                        cells.append(' '.join([item.text for item in self.iter_cell_items(cell)]))

                                    lines.append('\t'.join(cells))

                        new_file['text'] = '\n'.join(lines)
                    # Microsoft Excel workbooks
                    elif file_ext == '.xlsx':
                        lines = []
                        workbook = openpyxl.load_workbook(file_path, data_only = True)

                        for worksheet_name in workbook.sheetnames:
                            worksheet = workbook[worksheet_name]

                            for row in worksheet.rows:
                                cells = [
                                    (cell.value if cell.value is not None else '')
                                    for cell in row
                                ]

                                lines.append('\t'.join(cells))

                        new_file['text'] = '\n'.join(lines)

                    if self.main.settings_custom['file_area']['dialog_open_files']['auto_detect_langs']:
                        new_file['lang'] = wl_detection.detect_lang_text(self.main, new_file['text'])
                    else:
                        new_file['lang'] = self.main.settings_custom['files']['default_settings']['lang']

                    new_files.append(new_file)
                # Translation memory files
                else:
                    lines_src = []
                    lines_target = []

                    new_file_src = copy.deepcopy(new_file)
                    new_file_tgt = copy.deepcopy(new_file)

                    new_file_src['name'] = new_file_src['name_old'] = wl_checking_misc.check_new_name(f'{file_name}_source', file_names)
                    new_file_tgt['name'] = new_file_tgt['name_old'] = wl_checking_misc.check_new_name(f'{file_name}_target', file_names)

                    with open(file_path, 'r', encoding = new_file['encoding'], errors = 'replace') as f:
                        soup = bs4.BeautifulSoup(f.read(), 'lxml-xml')

                    # Extract source and target languages
                    elements_tuv = soup.select(r'tu:first-child tuv[xml\:lang]')

                    if len(elements_tuv) == 2:
                        new_file_src['lang'] = wl_conversion.to_iso_639_3(self.main, elements_tuv[0]['xml:lang'])
                        new_file_tgt['lang'] = wl_conversion.to_iso_639_3(self.main, elements_tuv[1]['xml:lang'])
                    else:
                        new_file_src['lang'] = new_file_tgt['lang'] = self.main.settings_custom['files']['default_settings']['lang']

                    for elements_tu in soup.select('tu'):
                        seg_src, seg_target = elements_tu.select('seg')

                        lines_src.append(seg_src.get_text().replace(r'\n', ' ').strip())
                        lines_target.append(seg_target.get_text().replace(r'\n', ' ').strip())

                    new_file_src['path'] = wl_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}_source.txt'))
                    new_file_tgt['path'] = wl_checking_misc.check_new_path(os.path.join(default_dir, f'{file_name}_target.txt'))

                    new_file_src['text'] = '\n'.join(lines_src)
                    new_file_tgt['text'] = '\n'.join(lines_target)

                    new_files.append(new_file_src)
                    new_files.append(new_file_tgt)

            if self.file_paths:
                self.main.settings_custom['imp']['files']['default_path'] = wl_misc.get_normalized_dir(self.file_paths[0])
        except Exception:
            err_msg = traceback.format_exc()

        self.progress_updated.emit(self.tr('Updating table...'))
        self.worker_done.emit(err_msg, new_files)

    # Reference: https://github.com/python-openxml/python-docx/issues/276
    def iter_block_items(self, parent):
        """
        Yield each paragraph and table child within *parent*, in document order.
        Each returned value is an instance of either Table or Paragraph. *parent*
        would most commonly be a reference to a main Document object, but
        also works for a _Cell object, which itself can contain paragraphs and tables.
        """
        if isinstance(parent, Document):
            parent_elm = parent.element.body
        elif isinstance(parent, _Cell):
            parent_elm = parent._tc
        else:
            raise ValueError("something's not right")

        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    def iter_cell_items(self, parent):
        parent_elm = parent._tc

        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                table = Table(child, parent)

                for row in table.rows:
                    for cell in row.cells:
                        yield from self.iter_cell_items(cell)

    # Reference: https://github.com/python-openxml/python-docx/issues/40
    def iter_visual_cells(self, table):
        prior_tcs = []
        visual_cells = []

        for row in table.rows:
            visual_cells.append([])

            for cell in row.cells:
                this_tc = cell._tc

                if this_tc in prior_tcs:  # skip cells pointing to same `<w:tc>` element
                    continue
                else:
                    prior_tcs.append(this_tc)

                    visual_cells[-1].append(cell)

        return visual_cells

class Table_Open_Files(wl_tables.Wl_Table_Add_Ins_Del_Clr):
    def __init__(self, parent):
        super().__init__(
            parent = parent,
            headers = [
                parent.tr('Path'),
                parent.tr('Encoding'),
                parent.tr('Language'),
                parent.tr('Tokenized'),
                parent.tr('Tagged')
            ],
            col_edit = 2
        )

        self.files_to_open = []

        self.setItemDelegateForColumn(0, wl_item_delegates.Wl_Item_Delegate_Uneditable(self))
        self.setItemDelegateForColumn(1, wl_item_delegates.Wl_Item_Delegate_Combo_Box_Custom(self, wl_boxes.Wl_Combo_Box_Encoding))
        self.setItemDelegateForColumn(2, wl_item_delegates.Wl_Item_Delegate_Combo_Box_Custom(self, wl_boxes.Wl_Combo_Box_Lang))
        self.setItemDelegateForColumn(3, wl_item_delegates.Wl_Item_Delegate_Combo_Box_Custom(self, wl_boxes.Wl_Combo_Box_Yes_No))
        self.setItemDelegateForColumn(4, wl_item_delegates.Wl_Item_Delegate_Combo_Box_Custom(self, wl_boxes.Wl_Combo_Box_Yes_No))

        self.button_clr.disconnect()
        self.button_clr.clicked.connect(lambda: self.clr_table(remove_placeholders = True))

        self.clr_table()

    def item_changed(self, item):
        super().item_changed(item = item)

        self.files_to_open = []

        if not self.is_empty():
            for row in range(self.model().rowCount()):
                self.files_to_open.append(self.model().item(row, 0).file)

    def del_row(self):
        for row in self.get_selected_rows():
            file_path = self.files_to_open[row]['path']

            if os.path.exists(file_path):
                os.remove(file_path)

        super().del_row()

    def clr_table(self, num_headers = 1, remove_placeholders = False):
        # Remove placeholders for new paths
        if remove_placeholders:
            for file in self.files_to_open:
                if os.path.exists(file['path']):
                    os.remove(file['path'])

        super().clr_table(num_headers = num_headers)

    def update_table(self):
        files = self.files_to_open

        if files:
            self.clr_table(len(files))

            self.disable_updates()

            for i, file in enumerate(files):
                self.model().setItem(i, 0, QStandardItem(file['path_original']))
                self.model().setItem(i, 1, QStandardItem(wl_conversion.to_encoding_text(self.main, file['encoding'])))
                self.model().setItem(i, 2, QStandardItem(wl_conversion.to_lang_text(self.main, file['lang'])))
                self.model().setItem(i, 3, QStandardItem(file['tokenized']))
                self.model().setItem(i, 4, QStandardItem(file['tagged']))

                self.model().item(i, 0).file = file

            self.enable_updates()
        else:
            self.clr_table()

class Wl_Worker_Open_Files(wl_threading.Wl_Worker):
    worker_done = pyqtSignal(str, list)

    def run(self):
        err_msg = ''
        new_files = []

        try:
            len_files = len(self.files_to_open)
            # Regex for headers
            re_tags_header = wl_matching.get_re_tags_with_tokens(self.main, tag_type = 'header')

            for i, file in enumerate(self.files_to_open):
                self.progress_updated.emit(self.tr(f'Opening files... ({i + 1}/{len_files})'))

                # Remove header tags
                tags_header = []

                for _, _, tag_opening in self.main.settings_custom['tags']['tags_header']:
                    tags_header.append(tag_opening[1:-1])

                with open(file['path'], 'w', encoding = file['encoding']) as f:
                    text = file['text']

                    if file['tagged'] == 'Yes' and tags_header:
                        # Use regex here since BeautifulSoup will add tags including <html> and <body> to the text
                        # See: https://www.crummy.com/software/BeautifulSoup/bs4/doc/#differences-between-parsers
                        text = re.sub(re_tags_header, '', text)

                    f.write(text)

                # Process texts
                file['text'] = wl_texts.Wl_Text(self.main, file)

                new_files.append(file)
        except Exception:
            err_msg = traceback.format_exc()

        self.progress_updated.emit(self.tr('Updating table...'))
        self.worker_done.emit(err_msg, new_files)

class Dialog_Open_Files(wl_dialogs.Wl_Dialog):
    def __init__(self, main):
        super().__init__(
            main,
            title = main.tr('Open Files'),
            width = 800
        )

        self.table_files = Table_Open_Files(self)

        self.table_files.model().itemChanged.connect(self.table_files_changed)

        self.table_files.button_add.hide()
        self.table_files.button_ins.hide()

        self.table_files.button_add_files = QPushButton(self.tr('Add files...'), self)
        self.table_files.button_add_folder = QPushButton(self.tr('Add folder...'), self)

        self.table_files.button_add_files.clicked.connect(self.add_files)
        self.table_files.button_add_folder.clicked.connect(self.add_folder)

        layout_table = wl_layouts.Wl_Layout()
        layout_table.addWidget(self.table_files, 0, 0, 5, 1)
        layout_table.addWidget(self.table_files.button_add_files, 0, 1)
        layout_table.addWidget(self.table_files.button_add_folder, 1, 1)
        layout_table.addWidget(self.table_files.button_del, 2, 1)
        layout_table.addWidget(self.table_files.button_clr, 3, 1)

        layout_table.setRowStretch(4, 1)

        self.checkbox_auto_detect_encodings = QCheckBox(self.tr('Auto-detect encodings'), self)
        self.checkbox_auto_detect_langs = QCheckBox(self.tr('Auto-detect languages'), self)
        self.checkbox_include_files_in_subfolders = QCheckBox(self.tr('Include files in subfolders'), self)

        self.checkbox_auto_detect_encodings.stateChanged.connect(self.settings_changed)
        self.checkbox_auto_detect_langs.stateChanged.connect(self.settings_changed)
        self.checkbox_include_files_in_subfolders.stateChanged.connect(self.settings_changed)

        layout_checkboxes = wl_layouts.Wl_Layout()
        layout_checkboxes.addWidget(self.checkbox_auto_detect_encodings, 0, 0)
        layout_checkboxes.addWidget(self.checkbox_auto_detect_langs, 0, 1)
        layout_checkboxes.addWidget(self.checkbox_include_files_in_subfolders, 1, 0, 1, 2)

        self.button_restore_default_settings = wl_buttons.Wl_Button_Restore_Default_Settings(self, load_settings = self.load_settings)
        self.button_open = QPushButton(self.tr('Open'), self)
        self.button_cancel = QPushButton(self.tr('Cancel'), self)

        self.button_open.clicked.connect(self.accept)
        self.button_cancel.clicked.connect(self.reject)

        self.setLayout(wl_layouts.Wl_Layout())
        self.layout().addLayout(layout_table, 0, 0, 1, 4)
        self.layout().addLayout(layout_checkboxes, 1, 0, 1, 4)

        self.layout().addWidget(wl_layouts.Wl_Separator(self), 2, 0, 1, 4)

        self.layout().addWidget(self.button_restore_default_settings, 3, 0)
        self.layout().addWidget(self.button_open, 3, 2)
        self.layout().addWidget(self.button_cancel, 3, 3)

        self.layout().setColumnStretch(1, 1)

        self.load_settings()

    def accept(self):
        super().accept()

        self.main.wl_file_area.table_files._open_files(files_to_open = self.table_files.files_to_open)

    def reject(self):
        # Remove placeholders for new paths
        for file in self.table_files.files_to_open:
            if os.path.exists(file['path']):
                os.remove(file['path'])

        super().reject()

    def load_settings(self, defaults = False):
        if defaults:
            settings = copy.deepcopy(self.main.settings_default['file_area']['dialog_open_files'])
        else:
            settings = copy.deepcopy(self.main.settings_custom['file_area']['dialog_open_files'])

        self.checkbox_auto_detect_encodings.setChecked(settings['auto_detect_encodings'])
        self.checkbox_auto_detect_langs.setChecked(settings['auto_detect_langs'])
        self.checkbox_include_files_in_subfolders.setChecked(settings['include_files_in_subfolders'])

        self.table_files.model().itemChanged.emit(QStandardItem())
        self.settings_changed()

    def table_files_changed(self, item):
        if self.table_files.is_empty():
            self.button_open.setEnabled(False)
        else:
            self.button_open.setEnabled(True)

    def settings_changed(self):
        settings = self.main.settings_custom['file_area']['dialog_open_files']

        settings['auto_detect_encodings'] = self.checkbox_auto_detect_encodings.isChecked()
        settings['auto_detect_langs'] = self.checkbox_auto_detect_langs.isChecked()
        settings['include_files_in_subfolders'] = self.checkbox_include_files_in_subfolders.isChecked()

    def _add_files(self, file_paths):
        def update_gui(err_msg, new_files):
            if not err_msg:
                self.table_files.files_to_open.extend(new_files)

                self.table_files.update_table()

                if file_paths_empty or file_paths_unsupported or file_paths_dup:
                    dialog_err_files = wl_dialogs_errs.Wl_Dialog_Err_Files(self.main, title = self.tr('Error Adding Files'))

                    dialog_err_files.label_err.set_text(self.tr('''
                        <div>
                            An error occurred while adding files, so the following file(s) were not added to the table.
                        </div>
                    '''))
                    dialog_err_files.table_err_files.model().setRowCount(len(file_paths_empty) + len(file_paths_unsupported) + len(file_paths_dup))

                    dialog_err_files.table_err_files.disable_updates()

                    for i, file_path in enumerate(file_paths_empty + file_paths_unsupported + file_paths_dup):
                        if file_path in file_paths_empty:
                            dialog_err_files.table_err_files.model().setItem(
                                i, 0,
                                QStandardItem(self.tr('Empty File'))
                            )
                        elif file_path in file_paths_unsupported:
                            dialog_err_files.table_err_files.model().setItem(
                                i, 0,
                                QStandardItem(self.tr('Unsupported File Type'))
                            )
                        elif file_path in file_paths_dup:
                            dialog_err_files.table_err_files.model().setItem(
                                i, 0,
                                QStandardItem(self.tr('Duplicate File'))
                            )

                        dialog_err_files.table_err_files.model().setItem(
                            i, 1,
                            QStandardItem(file_path)
                        )

                    dialog_err_files.table_err_files.enable_updates()

                    dialog_err_files.open()
            else:
                wl_dialogs_errs.Wl_Dialog_Err_Fatal(self.main, err_msg).open()

                wl_msgs.wl_msg_fatal_error(self.main)

        dialog_progress = wl_dialogs_misc.Wl_Dialog_Progress(self.main, text = self.tr('Checking files...'))

        file_paths, file_paths_unsupported = wl_checking_files.check_file_paths_unsupported(self.main, file_paths)
        file_paths, file_paths_empty = wl_checking_files.check_file_paths_empty(self.main, file_paths)
        file_paths, file_paths_dup = wl_checking_files.check_file_paths_dup(
            self.main,
            new_file_paths = file_paths,
            file_paths = [
                file['path_original']
                for file in self.main.settings_custom['file_area']['files_open'] + self.table_files.files_to_open
            ]
        )

        wl_threading.Wl_Thread(Wl_Worker_Add_Files(
            self.main,
            dialog_progress = dialog_progress,
            update_gui = update_gui,
            file_paths = file_paths,
            table = self.table_files
        )).start_worker()

    def add_files(self):
        if os.path.exists(self.main.settings_custom['imp']['files']['default_path']):
            default_dir = self.main.settings_custom['imp']['files']['default_path']
        else:
            default_dir = self.main.settings_default['imp']['files']['default_path']

        file_paths = QFileDialog.getOpenFileNames(
            parent = self.main,
            caption = self.tr('Open Files'),
            directory = wl_checking_misc.check_dir(default_dir),
            filter = ';;'.join(self.main.settings_global['file_types']['files']),
            initialFilter = self.main.settings_global['file_types']['files'][-1]
        )[0]

        if file_paths:
            self._add_files(file_paths)

    def add_folder(self):
        file_paths = []

        file_dir = QFileDialog.getExistingDirectory(
            parent = self.main,
            caption = self.tr('Open Folder'),
            directory = self.main.settings_custom['imp']['files']['default_path']
        )

        if file_dir:
            if self.main.settings_custom['file_area']['dialog_open_files']['include_files_in_subfolders']:
                for dir_path, dir_names, file_names in os.walk(file_dir):
                    for file_name in file_names:
                        file_paths.append(os.path.join(dir_path, file_name))
            else:
                file_names = list(os.walk(file_dir))[0][2]

                for file_name in file_names:
                    file_paths.append(os.path.join(file_dir, file_name))

            self._add_files(file_paths)

class Wl_Table_Files(wl_tables.Wl_Table):
    def __init__(self, parent):
        super().__init__(
            parent,
            headers = [
                parent.tr('Name'),
                parent.tr('Path'),
                parent.tr('Encoding'),
                parent.tr('Language'),
                parent.tr('Tokenized'),
                parent.tr('Tagged')
            ],
            editable = True,
            drag_drop = True
        )

        self.file_area = parent

        self.setItemDelegateForColumn(1, wl_item_delegates.Wl_Item_Delegate_Uneditable(self))
        self.setItemDelegateForColumn(2, wl_item_delegates.Wl_Item_Delegate_Uneditable(self))
        self.setItemDelegateForColumn(3, wl_item_delegates.Wl_Item_Delegate_Uneditable(self))
        self.setItemDelegateForColumn(4, wl_item_delegates.Wl_Item_Delegate_Uneditable(self))
        self.setItemDelegateForColumn(5, wl_item_delegates.Wl_Item_Delegate_Uneditable(self))

        self.selectionModel().selectionChanged.connect(self.selection_changed)
        self.clicked.connect(self.item_clicked)

        # Menu
        self.main.action_file_open_files.triggered.connect(self.open_files)
        self.main.action_file_open_dir.triggered.connect(self.open_dir)
        self.main.action_file_reopen.triggered.connect(self.reopen)

        self.main.action_file_select_all.triggered.connect(self.select_all)
        self.main.action_file_deselect_all.triggered.connect(self.deselect_all)
        self.main.action_file_invert_selection.triggered.connect(self.invert_selection)

        self.main.action_file_close_selected.triggered.connect(self.close_selected)
        self.main.action_file_close_all.triggered.connect(self.close_all)

    def item_changed(self, item):
        super().item_changed(item)

        if not self.is_empty():
            # Record old file names that might be useful for other slots
            self.file_area.file_names_old = list(self.file_area.get_selected_file_names())

            # Check for duplicate file names
            for row in range(self.model().rowCount()):
                file = self.model().item(row, 0).wl_file
                file_name = self.model().item(row, 0).text()

                if file_name != file['name_old']:
                    if self.main.wl_file_area.find_file_by_name(file_name):
                        self.disable_updates()

                        self.model().item(row, 0).setText(file['name_old'])

                        self.enable_updates()

                        wl_msg_boxes.Wl_Msg_Box_Warning(
                            self.main,
                            title = self.tr('Duplicate File Names Found'),
                            text = self.tr('''
                                <div>There is already a file with the same name in the file area.</div>
                                <div>Please specify a different file name.</div>
                            ''')
                        ).exec_()

                        self.setCurrentIndex(self.model().index(row, 0))

                        self.closeEditor(self.findChild(QLineEdit), QAbstractItemDelegate.NoHint)
                        self.edit(self.model().index(row, 0))

                    break

            self.main.settings_custom['file_area']['files_open'].clear()

            for row in range(self.model().rowCount()):
                file = self.model().item(row, 0).wl_file

                file['selected'] = True if self.model().item(row, 0).checkState() == Qt.Checked else False
                file['name'] = file['name_old'] = self.model().item(row, 0).text()
                file['encoding'] = wl_conversion.to_encoding_code(self.main, self.model().item(row, 2).text())
                file['lang'] = wl_conversion.to_lang_code(self.main, self.model().item(row, 3).text())
                file['tokenized'] = self.model().item(row, 4).text()
                file['tagged'] = self.model().item(row, 5).text()

                self.main.settings_custom['file_area']['files_open'].append(file)

        # Menu
        if not self.is_empty():
            self.main.action_file_select_all.setEnabled(True)
            self.main.action_file_deselect_all.setEnabled(True)
            self.main.action_file_invert_selection.setEnabled(True)

            self.main.action_file_close_all.setEnabled(True)
        else:
            self.main.action_file_select_all.setEnabled(False)
            self.main.action_file_deselect_all.setEnabled(False)
            self.main.action_file_invert_selection.setEnabled(False)

            self.main.action_file_close_all.setEnabled(False)

        if self.main.settings_custom['file_area']['files_closed']:
            self.main.action_file_reopen.setEnabled(True)
        else:
            self.main.action_file_reopen.setEnabled(False)

        self.selection_changed()

    def item_clicked(self):
        if not self.is_empty():
            for row in range(self.model().rowCount()):
                self.main.settings_custom['file_area']['files_open'][row]['selected'] = True if self.model().item(row, 0).checkState() == Qt.Checked else False

    def selection_changed(self):
        if self.get_selected_rows():
            self.main.action_file_close_selected.setEnabled(True)
        else:
            self.main.action_file_close_selected.setEnabled(False)

    def update_table(self):
        files = self.main.settings_custom['file_area']['files_open']

        if files:
            self.clr_table(len(files))

            self.disable_updates()

            for i, file in enumerate(files):
                item_name = QStandardItem(file['name'])
                # Record file properties
                item_name.wl_file = file
                item_name.setCheckable(True)

                if file['selected']:
                    item_name.setCheckState(Qt.Checked)
                else:
                    item_name.setCheckState(Qt.Unchecked)

                self.model().setItem(i, 0, item_name)
                self.model().setItem(i, 1, QStandardItem(file['path_original']))
                self.model().setItem(i, 2, QStandardItem(wl_conversion.to_encoding_text(self.main, file['encoding'])))
                self.model().setItem(i, 3, QStandardItem(wl_conversion.to_lang_text(self.main, file['lang'])))
                self.model().setItem(i, 4, QStandardItem(file['tokenized']))
                self.model().setItem(i, 5, QStandardItem(file['tagged']))

            self.enable_updates()
        else:
            self.clr_table(1)

    @wl_misc.log_timing
    def _open_files(self, files_to_open):
        def update_gui(err_msg, new_files):
            if not err_msg:
                len_files_old = len(self.main.settings_custom['file_area']['files_open'])

                self.main.settings_custom['file_area']['files_open'].extend(new_files)
                self.update_table()

                len_files_new = len(self.main.settings_custom['file_area']['files_open'])

                if len_files_new - len_files_old == 0:
                    self.main.statusBar().showMessage(self.tr('No files are newly opened!'))
                elif len_files_new - len_files_old == 1:
                    self.main.statusBar().showMessage(self.tr('1 file has been successfully opened.'))
                else:
                    self.main.statusBar().showMessage(self.tr(f'{len_files_new - len_files_old} files have been successfully opened.'))
            else:
                wl_dialogs_errs.Wl_Dialog_Err_Fatal(self.main, err_msg).open()

                wl_msgs.wl_msg_fatal_error(self.main)

        dialog_progress = wl_dialogs_misc.Wl_Dialog_Progress(self.main, text = self.tr('Checking files...'))

        wl_threading.Wl_Thread(Wl_Worker_Open_Files(
            self.main,
            dialog_progress = dialog_progress,
            update_gui = update_gui,
            files_to_open = files_to_open
        )).start_worker()

    def open_files(self):
        self.dialog_open_files = Dialog_Open_Files(self.main)
        self.dialog_open_files.open()

        self.dialog_open_files.table_files.button_add_files.click()

    def open_dir(self):
        self.dialog_open_files = Dialog_Open_Files(self.main)
        self.dialog_open_files.open()

        self.dialog_open_files.table_files.button_add_folder.click()

    def reopen(self):
        files = self.main.settings_custom['file_area']['files_closed'].pop()

        dialog_open_files = Dialog_Open_Files(self.main)
        dialog_open_files._add_files(list(dict.fromkeys([file['path_original'] for file in files])))

        self._open_files(files_to_open = dialog_open_files.table_files.files_to_open)

    def select_all(self):
        if self.model().item(0, 0):
            for i in range(self.model().rowCount()):
                if self.model().item(i, 0).checkState() == Qt.Unchecked:
                    self.model().item(i, 0).setCheckState(Qt.Checked)

    def deselect_all(self):
        if self.model().item(0, 0):
            for i in range(self.model().rowCount()):
                if self.model().item(i, 0).checkState() == Qt.Checked:
                    self.model().item(i, 0).setCheckState(Qt.Unchecked)

    def invert_selection(self):
        if self.model().item(0, 0):
            for i in range(self.model().rowCount()):
                if self.model().item(i, 0).checkState() == Qt.Checked:
                    self.model().item(i, 0).setCheckState(Qt.Unchecked)
                else:
                    self.model().item(i, 0).setCheckState(Qt.Checked)

    def _close_files(self, i_files):
        self.main.settings_custom['file_area']['files_closed'].append([])

        for i in reversed(i_files):
            file_to_remove = self.main.settings_custom['file_area']['files_open'].pop(i)

            self.main.settings_custom['file_area']['files_closed'][-1].append(file_to_remove)

            # Remove temporary files
            if os.path.exists(file_to_remove['path']):
                os.remove(file_to_remove['path'])

        self.update_table()

    def close_selected(self):
        self._close_files(self.get_selected_rows())

    def close_all(self):
        self._close_files(list(range(len(self.main.settings_custom['file_area']['files_open']))))

class Wrapper_File_Area(wl_layouts.Wl_Wrapper_File_Area):
    def __init__(self, main):
        super().__init__(main)

        self.file_names_old = []

        # Table
        self.table_files = Wl_Table_Files(self)

        self.wrapper_table.layout().addWidget(self.table_files, 0, 0)

        # Load files
        self.table_files.update_table()

    def get_selected_files(self):
        return (
            file
            for file in self.main.settings_custom['file_area']['files_open']
            if file['selected']
        )

    def get_selected_file_names(self):
        return (
            file['name']
            for file in self.get_selected_files()
        )

    def find_file_by_name(self, file_name, selected_only = False):
        if selected_only:
            files = self.get_selected_files()
        else:
            files = self.main.settings_custom['file_area']['files_open']

        for file in files:
            if file['name'] == file_name:
                return file

        return None

    def find_files_by_name(self, file_names, selected_only = False):
        files = [
            self.find_file_by_name(file_name, selected_only = selected_only)
            for file_name in file_names
        ]

        return (file for file in files if file)
